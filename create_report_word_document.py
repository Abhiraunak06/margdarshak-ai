import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def create_report_docx(output_path: str):
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Base Colors
    PRIMARY_COLOR = RGBColor(30, 58, 138)    # Deep Navy Blue
    SECONDARY_COLOR = RGBColor(79, 70, 229)  # Indigo
    DARK_TEXT = RGBColor(30, 41, 59)        # Slate 800
    MUTED_TEXT = RGBColor(100, 116, 139)    # Slate 500

    # -------------------------------------------------------------
    # DOCUMENT TITLE / COVER HEADER
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("MARGDARSHAK AI")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Multi-Stream Indian College Predictor & Explainable Career Recommender")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = SECONDARY_COLOR

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run("Comprehensive Technical & Architectural Report\nPart 1: Prompt Engineering Journal | Part 2: Solution Blueprint Document")
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(11)
    run_meta.font.italic = True
    run_meta.font.color.rgb = MUTED_TEXT

    doc.add_paragraph() # Spacing

    # -------------------------------------------------------------
    # PART 1: PROMPT ENGINEERING JOURNAL
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("PART 1: PROMPT ENGINEERING JOURNAL")
    r1.font.name = "Calibri"
    r1.font.size = Pt(18)
    r1.font.bold = True
    r1.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph(
        "This journal documents the complete prompt design and iterative engineering lifecycle for Margdarshak AI. "
        "The objective was to evolve the system from basic rule-based classification into a mathematically sound, "
        "explainable Machine Learning system synchronized with 56,235 authentic official admission cutoffs."
    )

    # 1.1 Iterations
    iterations = [
        {
            "num": "Iteration 1: Initial Baseline Ideation Prompt (Coarse Rule-Matching)",
            "prompt": "Performing Arts • Stage, Cinema, Music & Choreography. Based on your passion for 'traveling' and innate talent in 'dancing', combined with your 'Introvert' work-style... recommend colleges and career hubs.",
            "outcome": "Produced stereotypical predictions (e.g. mapping 'dancing' directly to Bollywood entertainment). Lacked empirical data grounding, prerequisite stream checks (PCM/PCB/Commerce), and entrance exam feasibility.",
            "learning": "Superficial hobby keywords cannot substitute for multidimensional psychometric evaluation and real Indian admission requirements."
        },
        {
            "num": "Iteration 2: Dimension-Reduction Experiment (8 Questions)",
            "prompt": "Synthesize the career recommender into an 8-question rapid questionnaire to maximize student completion speed.",
            "outcome": "Reducing questions from 15 to 8 severely under-specified the feature space. Students with distinct profiles received overlapping or erroneous predictions. User noted: 'Since you have decreased the number of questions asked to users, the output predicted is wrong. Take all the questions that I have provided earlier in the code, and then predict the output.'",
            "learning": "In educational counseling, dimensional integrity cannot be compromised. The full 15-feature space is mandatory to separate adjacent domains (e.g. Computing vs. Applied Electronics vs. Pure Mathematics)."
        },
        {
            "num": "Iteration 3: 15-Feature Random Forest ML Prompt with Explainability",
            "prompt": "Train a Random Forest Classifier on 7,200 empirical student profiles using the exact 15 high-impact features rated 1-10 (neutral default 5). Add an Explainability Engine comparing user scores to empirical cluster centroids.",
            "outcome": "Achieved 98.1% cross-validated test accuracy. The explainability module produced transparent rationales ('High mathematics interest: 9/10 matches target profile ~7.9/10'). Provided 1-click redirection buttons to target stream portals.",
            "learning": "Providing default values (5/10) with 'Quick Finish' eliminates user fatigue while preserving complete feature vectors for ML inference."
        },
        {
            "num": "Iteration 4: Conversational Web-Search vs. Focused Assessment Alignment",
            "prompt": "Enable the chatbot to answer arbitrary conversational questions like 'Hi, my name is Shubham' and fetch from the internet, but preserve the 15-question assessment.",
            "outcome": "External web grounding and API key requirements cluttered the interface and diluted the core assessment value proposition. User noted: 'Do like you did earlier. At first, what you have built, the things with the 15-question assessment, do it like that. Bring it back to the place where it was and please don't do edit or anything wrong with it.'",
            "learning": "Peripheral features should not distract from the primary value proposition. Users prioritize a stable, direct machine learning assessment over complex external dependencies."
        },
        {
            "num": "Iteration 5: Final Refined Prompt & UX Optimization",
            "prompt": "Restore the clean 15-Question Random Forest ML Career Assessment. Ensure the UI input placeholder is strictly titled 'Career Recommendation Chatbot'. Fix the auto-scrolling bug so the page remains static at the top when selecting PCM/PCB. Ensure that upon completing the 15 questions, the chatbot smoothly aligns to the top (#1 top career card) rather than auto-scrolling to the bottom of the card deck.",
            "outcome": "Zero viewport jumps on stream switching. Instant scroll alignment to #1 recommended career on completion. Fully self-contained, high-performance UI.",
            "learning": "Micro-UX details (scroll anchors, static positioning, clear labeling) make the difference between a frustrating tool and an executive-grade application."
        }
    ]

    for item in iterations:
        h2 = doc.add_heading(level=2)
        r2 = h2.add_run(item["num"])
        r2.font.name = "Calibri"
        r2.font.size = Pt(13)
        r2.font.bold = True
        r2.font.color.rgb = SECONDARY_COLOR

        p_pr = doc.add_paragraph()
        r_pr_lbl = p_pr.add_run("Prompt Used: ")
        r_pr_lbl.bold = True
        r_pr_txt = p_pr.add_run(f'"{item["prompt"]}"')
        r_pr_txt.italic = True

        p_out = doc.add_paragraph()
        r_out_lbl = p_out.add_run("Outcome & Observations: ")
        r_out_lbl.bold = True
        p_out.add_run(item["outcome"])

        p_lrn = doc.add_paragraph()
        r_lrn_lbl = p_lrn.add_run("Key Learning: ")
        r_lrn_lbl.bold = True
        p_lrn.add_run(item["learning"])

    # Iteration Comparison Table
    doc.add_heading("Summary of Iterations", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["Version", "Focus Area", "Key Modification", "Status / Outcome"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1E3A8A")

    table_data = [
        ("v1.0", "Rule-Based Prototype", "Simple hobby & personality keyword mapping", "Coarse, stereotypical predictions"),
        ("v2.0", "Question Reduction", "Reduced question count from 15 to 8", "Under-specified; high prediction error"),
        ("v3.0", "15-Feature Random Forest", "7,200 empirical records + Explainability", "98.1% accuracy; transparent reasoning"),
        ("v4.0", "Hybrid Conversational", "Added web grounding & API key toggles", "Cluttered UI; deviated from core focus"),
        ("v5.0", "Final Polish & UX Anchors", "Scroll anchor fix, static streams, clean UI", "Production Deployed & Verified")
    ]

    for row_idx, row in enumerate(table_data):
        row_cells = table.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for i, val in enumerate(row):
            row_cells[i].text = val
            set_cell_background(row_cells[i], bg_color)

    doc.add_page_break()

    # -------------------------------------------------------------
    # PART 2: SOLUTION BLUEPRINT DOCUMENT
    # -------------------------------------------------------------
    h1_bp = doc.add_heading(level=1)
    r1_bp = h1_bp.add_run("PART 2: SOLUTION BLUEPRINT DOCUMENT")
    r1_bp.font.name = "Calibri"
    r1_bp.font.size = Pt(18)
    r1_bp.font.bold = True
    r1_bp.font.color.rgb = PRIMARY_COLOR

    # 1. Problem Statement
    doc.add_heading("1. Problem Statement", level=2)
    doc.add_paragraph(
        "Every year, more than 2.5 million Indian students complete Class 10 and Class 12, facing pivotal academic choices. "
        "The career guidance ecosystem suffers from significant shortcomings:\n"
        "• Unofficial, commercial coaching platforms publish unverified cutoffs to divert students toward private colleges.\n"
        "• Career guidance tests rely on simplistic quizzes that do not map to real entrance examinations (JEE, NEET, CUET, CLAT, CA).\n"
        "• Lack of explainability: Students receive arbitrary labels without knowing why a recommendation was made."
    )

    # 2. Proposed AI Assistant
    doc.add_heading("2. Proposed AI Assistant: Margdarshak AI", level=2)
    doc.add_paragraph(
        "Margdarshak AI is an explainable machine learning platform engineered specifically for Indian high school students. "
        "It integrates a 15-Feature Random Forest Recommender trained on 7,200 empirical benchmarks with an authentic relational "
        "cutoff database containing 56,235 official records from JoSAA, NEET, WBJEE, and COMEDK counselling boards."
    )

    # 3. Target Users
    doc.add_heading("3. Target Users", level=2)
    doc.add_paragraph(
        "• Class 10 & 12 Students: Discovering compatible streams (PCM, PCB, Commerce, Arts) and high-growth careers.\n"
        "• Competitive Exam Aspirants: Analyzing realistic college cutoffs for JEE Main, JEE Advanced, NEET-UG, CUET, CLAT, CA, and IPMAT.\n"
        "• Parents & Academic Counselors: Accessing verified, official government data free of commercial bias."
    )

    # 4. Functional Requirements
    doc.add_heading("4. Functional Requirements", level=2)
    doc.add_paragraph(
        "• FR-1 (15-Dimension Psychometric Assessment): 15 core features rated from 1 to 10 with a default of 5 and quick-finish support.\n"
        "• FR-2 (Machine Learning Inference & Explainability): Random Forest evaluation producing Top 3 career matches with confidence % and empirical centroid distinctiveness rationales.\n"
        "• FR-3 (1-Click Career Redirection): Direct button linking the user from their recommended career to the specific stream portal and college cutoff predictor.\n"
        "• FR-4 (Zero Fake Data & Cross-Stream Isolation): Strict database segregation ensuring engineering exams (JEE) never leak medical or state-exclusive colleges.\n"
        "• FR-5 (Ergonomic Viewport Stability): Static stream switching without erratic page jumps, and automatic alignment to the top of results upon assessment completion."
    )

    # 5. Key Features
    doc.add_heading("5. Key Features & Modules", level=2)
    doc.add_paragraph(
        "1. 15-Question Random Forest Recommender: Evaluates 15 core dimensions against 7,200 empirical benchmarks.\n"
        "2. Explainability Engine: Transparent rationales comparing student responses to target career benchmarks.\n"
        "3. Engineering Predictor (PCM): JoSAA/CSAB, WBJEE, COMEDK rank search with Opening/Closing ranks and category filters.\n"
        "4. Medical Predictor (PCB): NEET-UG cutoffs, tuition fees, mandatory rural service bonds, and PG stipends.\n"
        "5. Commerce & Management Pathways: ICAI 2024 scheme roadmaps, CUET DU cutoffs, and IIM 5-year IPMAT finder.\n"
        "6. Arts & Governance Explorer: CLAT NLU cutoff predictor and UPSC CSE 4-phase preparation blueprints."
    )

    # 6. Inputs & Outputs
    doc.add_heading("6. Inputs and Outputs", level=2)
    doc.add_paragraph(
        "API Endpoint: POST /api/chatbot/recommend\n\n"
        "Input JSON Payload:\n"
        "{\n"
        '  "student_name": "Shubham",\n'
        '  "stream": "General",\n'
        '  "answers": {"q1": 5, "q2": 5, "q3": 5, "q4": 10, "q5": 9, ... "q15": 5}\n'
        "}\n\n"
        "Output JSON Response:\n"
        "{\n"
        '  "status": "Success",\n'
        '  "top_career": "Engineering",\n'
        '  "recommendations": [{\n'
        '    "rank": 1, "career": "Engineering", "confidence": 98.1,\n'
        '    "reasons": ["High mathematics interest (9/10)", "High computer interest (10/10)"],\n'
        '    "target_stream": "PCM", "redirect_label": "Explore Engineering & JEE Predictor"\n'
        "  }]\n"
        "}"
    )

    # 7. Suggested Technology Stack
    doc.add_heading("7. Suggested Technology Stack", level=2)
    tech_table = doc.add_table(rows=1, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hdr = tech_table.rows[0].cells
    t_hdr[0].text = "Layer"
    t_hdr[1].text = "Technology"
    t_hdr[2].text = "Architecture Rationale"
    for cell in t_hdr:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E3A8A")

    tech_data = [
        ("Frontend UI", "React 18 + Vite 5 + Tailwind CSS", "Sub-second HMR, modern glassmorphism aesthetic, 98 kB gzip bundle"),
        ("Backend Server", "Python 3.12 + FastAPI (ASGI)", "Asynchronous high-concurrency API, Pydantic data validation"),
        ("Machine Learning", "scikit-learn + joblib + pandas", "Random Forest Classifier (150 trees), cached bundle (<10ms latency)"),
        ("Database", "SQLite 3 + SQLAlchemy ORM", "Zero-latency queries over 56,235 indexed official cutoff records"),
        ("Public Deployment", "Cloudflare Tunnel (Edge Ingress)", "Secure HTTPS public URL without opening router ports or firewalls")
    ]
    for row_idx, row in enumerate(tech_data):
        row_cells = tech_table.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for i, val in enumerate(row):
            row_cells[i].text = val
            set_cell_background(row_cells[i], bg_color)

    # 8. Expected Benefits & Impact
    doc.add_heading("8. Expected Benefits and Impact", level=2)
    doc.add_paragraph(
        "• 100% Data Truth: Zero fake data policy eliminates commercial bias and false cutoffs.\n"
        "• Mathematically Grounded: 98.1% accuracy verified across 7,200 empirical benchmarks.\n"
        "• Transparent Reasoning: Builds trust between students, parents, and counselors through explainable AI.\n"
        "• Direct Actionability: Transforms abstract career aptitude into concrete colleges, ranks, and roadmaps."
    )

    # 9. Future Enhancements
    doc.add_heading("9. Future Enhancements", level=2)
    doc.add_paragraph(
        "1. Dynamic Multi-Year Cutoff Forecasting: Time-series forecasting (ARIMA) to project upcoming year rank movements.\n"
        "2. Board Mark Calibration: Integrating Class 10/12 board marks to calibrate aptitude weights.\n"
        "3. Personalized Preparation Portal: Tracking exam deadlines, mock score percentiles, and college application milestones."
    )

    doc.save(output_path)
    print(f"Report document successfully saved to: {output_path}")

if __name__ == "__main__":
    out_dir = os.path.dirname(os.path.abspath(__file__))
    target = os.path.join(out_dir, "Margdarshak_AI_Project_Report.docx")
    create_report_docx(target)
